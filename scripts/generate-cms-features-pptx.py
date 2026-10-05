#!/usr/bin/env python3
"""Generate Spectra feature overview PowerPoint."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x14, 0x24)
NAVY_2 = RGBColor(0x12, 0x1E, 0x33)
CARD = RGBColor(0x17, 0x27, 0x3F)
CARD_LINE = RGBColor(0x2A, 0x3F, 0x5C)
ACCENT = RGBColor(0x3B, 0x82, 0xF6)
ACCENT_2 = RGBColor(0x22, 0xD3, 0xEE)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
SOFT = RGBColor(0xCB, 0xD5, 0xE1)

W = Inches(13.333)
H = Inches(7.5)
LOGO = "/Users/anshikatrivedi/Digital-Signage-Orion-1/docs/spectra-logo.png"


def set_run(run, text, size=18, bold=False, color=WHITE, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text_box(slide, l, t, w, h, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        set_run(p.add_run(), line, size, bold, color, font)
    return box


def set_multiline(tf, text, size=13, color=SOFT, first_size=None):
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2 if line.strip() else 6)
        display = line if line.strip() else " "
        set_run(p.add_run(), display, first_size if i == 0 and first_size else size, False, color)


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def card(slide, l, t, w, h, fill=CARD):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill_shape(s, fill)
    s.line.color.rgb = CARD_LINE
    s.line.width = Pt(1)
    s.adjustments[0] = 0.08
    return s


def accent_bar(slide, l, t, w=Inches(0.08), h=Inches(0.42)):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill_shape(s, ACCENT)
    return s


def bg(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill_shape(s, NAVY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.06))
    fill_shape(line, ACCENT)
    logo_mark(slide, Inches(10.85), Inches(0.16), Inches(1.85))


def footer(slide, page, total=11):
    add_text_box(slide, Inches(0.5), Inches(7.12), Inches(8), Inches(0.28), "Spectra  ·  Feature overview", 11, False, MUTED)
    add_text_box(slide, Inches(11.2), Inches(7.12), Inches(1.6), Inches(0.28), f"{page}  /  {total}", 11, False, MUTED, PP_ALIGN.RIGHT)


def logo_mark(slide, left, top, width):
    height = width * (503 / 916)
    plate = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left - Inches(0.08),
        top - Inches(0.06),
        width + Inches(0.16),
        height + Inches(0.12),
    )
    fill_shape(plate, WHITE)
    plate.line.fill.background()
    plate.adjustments[0] = 0.15
    slide.shapes.add_picture(LOGO, left, top, width, height)


def kicker(slide, text):
    add_text_box(slide, Inches(0.55), Inches(0.22), Inches(8), Inches(0.28), text.upper(), 11, True, ACCENT_2)


def title(slide, text, t=Inches(0.42)):
    add_text_box(slide, Inches(0.5), t, Inches(10.1), Inches(0.55), text, 28, True, WHITE)


def subtitle(slide, text, t=Inches(0.95)):
    add_text_box(slide, Inches(0.5), t, Inches(10.1), Inches(0.4), text, 15, False, MUTED)


def feature_card(slide, l, t, w, h, heading, body, tag=None):
    card(slide, l, t, w, h)
    accent_bar(slide, l + Inches(0.18), t + Inches(0.22), Inches(0.07), Inches(0.28))
    add_text_box(slide, l + Inches(0.35), t + Inches(0.16), w - Inches(0.5), Inches(0.38), heading, 16, True, WHITE)
    box = slide.shapes.add_textbox(l + Inches(0.22), t + Inches(0.54), w - Inches(0.4), h - Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    set_multiline(tf, body)
    if tag:
        add_text_box(slide, l + Inches(0.22), t + h - Inches(0.38), w - Inches(0.4), Inches(0.28), tag, 11, True, ACCENT_2)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    total = 11

    # 1 Title
    s = prs.slides.add_slide(blank)
    sheet = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill_shape(sheet, WHITE)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.08))
    fill_shape(bar, ACCENT)
    logo_w = Inches(6.4)
    logo_h = logo_w * (503 / 916)
    s.shapes.add_picture(LOGO, (W - logo_w) / 2, Inches(1.35), logo_w, logo_h)
    add_text_box(s, Inches(0.7), Inches(4.85), Inches(12), Inches(0.7), "Feature overview", 36, True, NAVY, PP_ALIGN.CENTER)
    add_text_box(
        s,
        Inches(1.4),
        Inches(5.6),
        Inches(10.5),
        Inches(0.9),
        "What is live in the client workspace, the internal platform portal,\nand how both drive Android players in the field.",
        18,
        False,
        RGBColor(0x47, 0x55, 0x69),
        PP_ALIGN.CENTER,
    )
    add_text_box(s, Inches(0.7), Inches(6.85), Inches(12), Inches(0.3), "Internal product briefing", 13, False, MUTED, PP_ALIGN.CENTER)

    # 2 Two portals
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "Architecture")
    title(s, "Two portals, one product")
    subtitle(s, "Client operators run screens. Internal staff run tenants. They never share a workspace.")
    feature_card(
        s, Inches(0.5), Inches(1.55), Inches(6.0), Inches(4.85),
        "Client workspace  ·  /app",
        "Login as Client dashboard.\n\nContent, playlists, schedules, tickers, screens, India map, and proof of play.\n\nRoles: Org Admin, Manager, Content Editor, Analyst Viewer — with VIEW / EDIT / MANAGE / CONTROL per feature.",
        "Tenant-facing signage operations",
    )
    feature_card(
        s, Inches(6.8), Inches(1.55), Inches(6.0), Inches(4.85),
        "Platform portal  ·  /platform",
        "Login as Platform portal.\n\nCreate organizations, set device limits, invite the first org admin, manage internal team.\n\nRoles: Super Admin, Platform Admin, Sales, Support.",
        "Internal operations",
    )
    footer(s, 2, total)

    # 3 Pipeline
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "How it fits together")
    title(s, "Content to screen to proof")
    subtitle(s, "Spectra is the control plane. The Android player is the runtime.")
    steps = [
        ("1", "Library", "Upload media into folders"),
        ("2", "Playlists", "Order assets & durations"),
        ("3", "Schedule / ticker", "Time takeovers & overlays"),
        ("4", "Screens", "Pair, locate, configure"),
        ("5", "Player", "Offline loop + sync"),
        ("6", "Insights", "Proof of play logs"),
    ]
    x = Inches(0.45)
    for num, head, body in steps:
        card(s, x, Inches(1.7), Inches(2.0), Inches(4.4))
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.7), Inches(2.05), Inches(0.55), Inches(0.55))
        fill_shape(circ, ACCENT)
        add_text_box(s, x + Inches(0.7), Inches(2.12), Inches(0.55), Inches(0.45), num, 16, True, WHITE, PP_ALIGN.CENTER)
        add_text_box(s, x + Inches(0.12), Inches(2.8), Inches(1.76), Inches(0.7), head, 16, True, WHITE, PP_ALIGN.CENTER)
        add_text_box(s, x + Inches(0.12), Inches(3.5), Inches(1.76), Inches(1.6), body, 13, False, SOFT, PP_ALIGN.CENTER)
        x += Inches(2.14)
    footer(s, 3, total)

    # 4 Overview
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "Spectra  ·  Overview")
    title(s, "Dashboard at a glance")
    items = [
        ("Fleet health", "Online, warning, and offline counts for paired screens."),
        ("Playback trend", "Plays today and over the last 7 days, plus top screens."),
        ("Library mix", "Share of video, image, web, and document assets."),
        ("Schedule preview", "What is live or coming up on the calendar."),
        ("Screen Locations", "India map of paired devices with valid install coordinates."),
        ("Activity", "Recent operator and device events."),
    ]
    positions = [
        (0.5, 1.55), (4.55, 1.55), (8.6, 1.55),
        (0.5, 4.15), (4.55, 4.15), (8.6, 4.15),
    ]
    for (l, t), (head, body) in zip(positions, items):
        feature_card(s, Inches(l), Inches(t), Inches(3.85), Inches(2.35), head, body)
    footer(s, 4, total)

    # 5 Screens
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "Spectra  ·  Screens")
    title(s, "Device management")
    subtitle(s, "Pair Android players, then operate each screen from a details panel.")
    left = [
        "6-character pairing code → device token",
        "Spectra list shows paired devices only",
        "Status, last sync, cache, player version",
        "Assign or unassign a playlist",
        "Reboot, unregister, or delete",
    ]
    right = [
        "Installation location: search or drop a pin",
        "India-only search, save, and reverse geocode",
        "Landscape / portrait and Stretch to Fit",
        "Default durations: image, video, document, URL",
        "Targeted tickers for this screen only",
    ]
    card(s, Inches(0.5), Inches(1.55), Inches(6.05), Inches(4.85))
    add_text_box(s, Inches(0.75), Inches(1.75), Inches(5.5), Inches(0.4), "Fleet operations", 18, True, WHITE)
    box = s.shapes.add_textbox(Inches(0.75), Inches(2.25), Inches(5.5), Inches(3.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(left):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        set_run(p.add_run(), "▸  " + line, 15, False, SOFT)
    card(s, Inches(6.75), Inches(1.55), Inches(6.05), Inches(4.85))
    add_text_box(s, Inches(7.0), Inches(1.75), Inches(5.5), Inches(0.4), "Per-screen controls", 18, True, WHITE)
    box = s.shapes.add_textbox(Inches(7.0), Inches(2.25), Inches(5.5), Inches(3.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(right):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        set_run(p.add_run(), "▸  " + line, 15, False, SOFT)
    footer(s, 5, total)

    # 6 Playlists + Library
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "Spectra  ·  Content")
    title(s, "Library and playlists")
    feature_card(
        s, Inches(0.5), Inches(1.5), Inches(6.05), Inches(5.0),
        "Library",
        "Upload and organize media.\n\nTypes: image, video, document (PDF / Office), and URL.\n\nFolders are for Spectra organization. The player still receives a flat ordered list.\n\nReady assets get download URLs for offline cache.",
    )
    feature_card(
        s, Inches(6.75), Inches(1.5), Inches(6.05), Inches(5.0),
        "Playlists",
        "Create, rename, reorder, and assign to screens.\n\nEach slot can set an explicit duration, or stay blank to use that device’s default for the asset type.\n\nRule: playlist override always wins. Null means “ask the device.”\n\nCampaigns were removed — playlists point straight at assets.",
    )
    footer(s, 6, total)

    # 7 Tickers + Schedule
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "Spectra  ·  Overlays and time")
    title(s, "Tickers and scheduling")
    feature_card(
        s, Inches(0.5), Inches(1.5), Inches(6.05), Inches(5.0),
        "Tickers",
        "Scrolling overlay on playlist content.\n\n• Speed: slow / normal / fast\n• Style: classic, neon, gradient, minimal\n• Position: top or bottom\n• Height: 10–20% of screen\n• Priority and colors\n• Scope: all devices, or selected screens\n• Device details can add a screen-only ticker",
    )
    feature_card(
        s, Inches(6.75), Inches(1.5), Inches(6.05), Inches(5.0),
        "Scheduling",
        "Time-boxed playlist takeovers.\n\n• Start and end date-time with timezone\n• One device or every device\n• Enable / disable without deleting\n• Status is derived: scheduled, active, completed, disabled\n\nPlayers get nextContentChangeAt so a schedule start does not wait for the next heartbeat.",
    )
    footer(s, 7, total)

    # 8 Insights
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "Spectra  ·  Insights")
    title(s, "Proof of play reporting")
    subtitle(s, "Players queue one log per play and upload to POST /api/player/pop-logs.")
    bullets = [
        ("Filters", "Today, yesterday, last 7 / 15 days, or a custom range. Device, status, folder, search."),
        ("KPIs & charts", "Impressions, fidelity, engagement, device mix, top content."),
        ("Log table", "Device, playlist, asset, start, end, duration, verified / failed."),
        ("Excel export", "Same rows as the Spectra table. UTC stored, shown in the operator timezone."),
        ("Integrity", "Retries are de-duplicated. Clock-skewed future stamps are stored but flagged."),
        ("Locations map", "Dashboard markers stay on saved Indian coordinates. No fake pins."),
    ]
    positions = [
        (0.5, 1.5), (4.55, 1.5), (8.6, 1.5),
        (0.5, 4.15), (4.55, 4.15), (8.6, 4.15),
    ]
    for (l, t), (head, body) in zip(positions, bullets):
        feature_card(s, Inches(l), Inches(t), Inches(3.85), Inches(2.35), head, body)
    footer(s, 8, total)

    # 9 Locations
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "Spectra  ·  Screen Locations")
    title(s, "India-only installation map")
    subtitle(s, "Admin-assigned coordinates are the source of truth for the dashboard map.")
    points = [
        ("Search", "Nominatim with countrycodes=in. Phoenix Mall, Lucknow, Delhi — Indian results only."),
        ("Validate", "Save blocked outside India: “Please select a location in India.”"),
        ("Persist", "Latitude, longitude, address, city, state, country=India. Source stays ADMIN_ASSIGNED or GEOCODED."),
        ("Map", "Picker opens on India. Dashboard shows paired screens with valid coords. GPS cannot overwrite unless allowed."),
    ]
    y = Inches(1.5)
    for head, body in points:
        card(s, Inches(0.5), y, Inches(12.3), Inches(1.15))
        accent_bar(s, Inches(0.7), y + Inches(0.38), Inches(0.08), Inches(0.38))
        add_text_box(s, Inches(1.0), y + Inches(0.18), Inches(2.2), Inches(0.8), head, 18, True, WHITE)
        add_text_box(s, Inches(3.3), y + Inches(0.28), Inches(9.1), Inches(0.7), body, 15, False, SOFT)
        y += Inches(1.28)
    footer(s, 9, total)

    # 10 Settings
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "Spectra  ·  Settings")
    title(s, "Access and workspace control")
    feature_card(
        s, Inches(0.5), Inches(1.5), Inches(6.05), Inches(5.0),
        "Implemented",
        "Invite members by email.\n\nRoles: Org Admin, Manager, Content Editor, Analyst Viewer.\n\nPer-feature permissions: Dashboard, Assets, Playlists, Schedule, Tickers, Devices, Reports, Team, Settings.\n\nRead-only users see the same pages but cannot save.\n\nOrganization switcher for users in more than one tenant.",
    )
    feature_card(
        s, Inches(6.75), Inches(1.5), Inches(6.05), Inches(5.0),
        "UI only / limited",
        "Profile photo, theme, accent, and notification toggles are present in Settings chrome.\n\nSeveral of those controls are local UI (toast save) rather than a full account-settings API.\n\nTreat Access Management as the real settings product; the rest is presentation.",
        "Do not demo as fully persisted",
    )
    footer(s, 10, total)

    # 11 Platform
    s = prs.slides.add_slide(blank)
    bg(s)
    kicker(s, "Platform portal")
    title(s, "Internal command center")
    items = [
        ("Organizations", "Create, activate, suspend client workspaces. Set a per-org device limit."),
        ("First admin", "Invite the org admin before handing over Spectra."),
        ("Team", "Internal users, ownership, and support access."),
        ("Reminders", "Internal alerts for ops follow-up."),
        ("Billing", "Placeholder — labeled planned. Not a live subscription workflow."),
        ("Split", "Keep billing, support, and tenant setup out of the client workspace."),
    ]
    positions = [
        (0.5, 1.5), (4.55, 1.5), (8.6, 1.5),
        (0.5, 4.15), (4.55, 4.15), (8.6, 4.15),
    ]
    for (l, t), (head, body) in zip(positions, items):
        feature_card(s, Inches(l), Inches(t), Inches(3.85), Inches(2.35), head, body)
    footer(s, 11, total)

    out = "/Users/anshikatrivedi/Digital-Signage-Orion-1/docs/Spectra-Features.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
